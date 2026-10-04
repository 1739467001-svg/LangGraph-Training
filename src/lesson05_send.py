# -*- coding: utf-8 -*-
"""
第 5 课：并行分发（Send）—— 多个批次同时抽取

真实场景：测试患者甲有 200 条医嘱，一次全塞给模型，上下文装不下、速度也慢。
把医嘱切成 N 批，开 N 个「抽取工位」同时干活，最后把结果合并——
这就是 map（分发）→ reduce（合并）。LangGraph 里的关键字叫 Send。

本课两个新概念：

1. **Send**：条件边的返回值不只能是「下一个节点名」，还可以是一串
   Send(节点名, 子状态)。LangGraph 会为每个 Send 开一个该节点的独立实例，
   并行执行——图第一次长出「分叉」，而且分几叉是运行时按数据量决定的。

2. **reducer（归并规则）**：N 个并行节点同时往袋子的同一个格子写，
   必须规定「怎么叠」：
       facts: Annotated[list, operator.add]   ← 追加（本课要用）
       facts: list                             ← 覆盖（并行时当场打架报错）

图：
    START → load_orders ──(Send×N 并行)──→ extract_chunk ──→ merge_draft → END
                          分几叉由医嘱量决定

⚠ 预告一个故意的体验：先填 TODO 1（真并行）再跑，会撞上并行写冲突的报错；
   填 TODO 2（reducer）治好它。痛一次才记得住。

运行：/usr/bin/python3 src/lesson05_send.py 测试患者甲
"""
import json
import math
import os
import re
import sys

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import Annotated, List, TypedDict
import operator
from pydantic import BaseModel, Field
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langchain_core.prompts import ChatPromptTemplate
from langgraph.graph import StateGraph, START, END
from langgraph.types import Send
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")
CHUNK_SIZE = 20  # 每个工位处理多少条医嘱


# ---------- 合同 / 替身模型 / 提示词（同第 3、4 课）----------
class MedicationFact(BaseModel):
    drug: str = Field(description="药品名称，逐字照抄医嘱原文")
    dose: str = Field(description="剂量+单位；原文找不到依据就填【待医生补充】")
    route: str = Field(description="给药途径")
    source: str = Field(description="出处，格式：医嘱.xlsx 第N行 医嘱号M")
    freq: str = Field(description="用药频次，逐字照抄医嘱原文；找不到依据就填【待医生补充】")


class FactList(BaseModel):
    facts: List[MedicationFact]


FIXED_SCRIPT = [
    {"drug": "0.9%氯化钠注射液", "dose": "100ml", "route": "静脉滴注",
     "freq": "每日1次", "source": "医嘱.xlsx 第20行 医嘱号10"},
    {"drug": "注射用头孢呋辛钠", "dose": "1.5g", "route": "静脉滴注",
     "freq": "每8小时1次", "source": "医嘱.xlsx 第9999行 医嘱号404"},
    {"drug": "阿司匹林肠溶片", "dose": "【待医生补充】", "route": "口服",
     "freq": "每日1次", "source": "医嘱.xlsx 第33行 医嘱号21"},
]


class FakeMedModel(BaseChatModel):
    @property
    def _llm_type(self) -> str:
        return "fake-med-model"

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        content = json.dumps({"facts": FIXED_SCRIPT}, ensure_ascii=False)
        return ChatResult(generations=[ChatGeneration(message=AIMessage(content=content))])


prompt = ChatPromptTemplate.from_messages([
    ("system", """你是一名住院病历的用药信息抽取助手。请从下面的医嘱列表里抽取用药事实。

规则：
1. 只使用医嘱原文里出现的信息，药名、剂量逐字照抄，不改写不翻译；
2. 每一条事实都必须给出处，格式：医嘱.xlsx 第N行 医嘱号M；
3. 原文里没有的信息，一律填【待医生补充】，绝不猜测、绝不编造数值；
4. 你只负责抽取事实，不判断用药对错、不下诊断、不建议换药——那是医生的工作。"""),
    ("human", "医嘱列表（共 {order_count} 条）：\n{orders_text}"),
])

extraction_chain = prompt | FakeMedModel() | (lambda m: FactList.model_validate_json(m.content))


# ---------- State：facts 带归并规则 ----------
class VisitState(TypedDict):
    patient_id: str
    orders: list
    # 参考答案（TODO 2）：归并规则 = 追加。
    # Annotated[list, operator.add] 告诉 LangGraph：这个格子有多个写入时，
    # 把所有列表拼接起来，而不是报错或互相覆盖
    facts: Annotated[list, operator.add]
    draft: str


def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        dose = " ".join(str(x) for x in (row[idx["DOSAGE"]], row[idx["DOSAGE_UNITS"]]) if x)
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]],
            "text": str(text), "dose": dose or "【待医生补充】",
            "route": row[idx["ADMINISTRATION"]] or "【待医生补充】",
            "freq": row[idx["FREQUENCY"]] or "【待医生补充】",
        })
    return {"orders": orders, "patient_id": patient_id}


# ---------- 分发：一个节点变 N 个工位 ----------
def fan_out(state: VisitState):
    """条件边的返回值是 Send 列表——返回几个，就并行开几个工位"""
    orders = state["orders"]
    chunks = [orders[i:i + CHUNK_SIZE] for i in range(0, len(orders), CHUNK_SIZE)]
    print(f"[分发] {len(orders)} 条医嘱 → 切成 {len(chunks)} 批（每批 {CHUNK_SIZE} 条）")
    # 参考答案（TODO 1）：返回 Send 列表——几个元素就并行开几个工位，
    # 第二个参数是该工位的「喂料」（它自己的输入状态）
    return [Send("extract_chunk", {"patient_id": state["patient_id"],
                                   "chunk_no": i, "orders": ch})
            for i, ch in enumerate(chunks)]


def extract_chunk(state: dict) -> dict:
    """一个工位：只处理分给自己的那一批。Send 传来的 dict 就是这里的输入状态"""
    chunk = state["orders"]
    no = state.get("chunk_no", "?")
    orders_text = "\n".join(
        f"第{o['row']}行 医嘱号{o['order_no']}：{o['text']} {o['dose']} {o['route']} {o['freq']}"
        for o in chunk
    )
    result: FactList = extraction_chain.invoke({
        "order_count": len(chunk), "orders_text": orders_text,
    })
    facts = [f.model_dump() for f in result.facts]
    print(f"[工位{no}] 收到 {len(chunk)} 条 → 交货 {len(facts)} 条")
    return {"facts": facts}


def merge_draft(state: VisitState) -> dict:
    facts = state["facts"]
    lines = [f"用药草稿（并行抽取）：{state['patient_id']}"
             f" —— {math.ceil(len(state['orders']) / CHUNK_SIZE)} 个工位共交货 {len(facts)} 条"]
    for i, f in enumerate(facts, start=1):
        lines.append(
            f"{i}. {f['drug']} | {f['dose']} | {f['route']} | 频次:{f['freq']}"
            f" | 出处：{f['source']}"
        )
    lines.append("")
    lines.append("⚠ 草稿由模型起草：每一条都必须核对后才能使用（第 4 课的核对节点在这里接上）。")
    return {"draft": "\n".join(lines)}


def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("extract_chunk", extract_chunk)
    g.add_node("merge_draft", merge_draft)

    g.add_edge(START, "load_orders")
    # 条件边的路由函数返回 Send 列表 → 动态并行
    g.add_conditional_edges("load_orders", fan_out)
    g.add_edge("extract_chunk", "merge_draft")
    g.add_edge("merge_draft", END)

    app = g.compile()
    final = app.invoke({"patient_id": patient_id, "orders": [], "facts": [], "draft": ""})

    print()
    print(final["draft"])

    out_dir = os.path.join(os.path.dirname(__file__), "..", "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"lesson05_{patient_id}.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(final["draft"] + "\n")
    print(f"[已保存] {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
