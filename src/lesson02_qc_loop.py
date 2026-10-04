# -*- coding: utf-8 -*-
"""
第 2 课：条件边与循环 —— 质控不通过就退回重写，最多 2 次

图结构（第一次出现「回头路」）：

    START → load_orders → build_timeline → qc_check ──通过──→ render_report → END
                                            ↑   │不通过（且没到重试上限）
                                            │   ↓
                                            └─ fix_issues（修一轮，再送回质控）

质控规则（本课先用代码规则做质控；第 3 课大模型才登场）：
  拦下（不通过）：剂量是 "?" —— 医嘱表里没填，也没法从文本确认
  提示（不拦下）：途径缺失、没有开始时间——记进报告但不退回
  防编造兜底：修了还修不出来的，明确标【待医生补充】，绝不编个数凑数

运行：
    /usr/bin/python3 src/lesson02_qc_loop.py
"""
import os
import re
import sys
from datetime import datetime

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import TypedDict
from langgraph.graph import StateGraph, START, END
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")

MAX_REWRITE = 2  # 最多退回重写次数


# ---------- State：比第 1 课多了三个格子 ----------
class VisitState(TypedDict):
    patient_id: str
    orders: list
    timeline: list
    report: str
    attempts: int    # 已经重写过几轮
    qc: dict         # 质控结论：{"passed": bool, "problems": [...], "warnings": [...]}
    notes: list      # 处理过程的备注（提取了多少、标记了多少）


# ---------- 节点 1、2：和第 1 课相同，直接搬过来 ----------
def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]],
            "text": str(text),
            "dosage": row[idx["DOSAGE"]],
            "units": row[idx["DOSAGE_UNITS"]],
            "route": row[idx["ADMINISTRATION"]],
            "start": row[idx["START_DATE_TIME"]],
        })
    return {"orders": orders, "patient_id": patient_id}


def build_timeline(state: VisitState) -> dict:
    orders = list(state["orders"])
    orders.sort(key=lambda o: o["start"] or datetime.max)
    timeline = []
    for i, o in enumerate(orders, start=1):
        dose = " ".join(str(x) for x in (o["dosage"], o["units"]) if x) or "?"
        timeline.append({
            "no": f"T{i:03d}", "time": o["start"], "text": o["text"],
            "dose": dose, "route": o["route"] or "?",
            "source": f"医嘱.xlsx 第{o['row']}行 医嘱号{o['order_no']}",
        })
    return {"timeline": timeline}


# ---------- 节点 3：质控（本课新角色）----------
def qc_check(state: VisitState) -> dict:
    """检查时间线，把问题分成两类：拦下的（problems）和提示的（warnings）"""
    problems, warnings = [], []
    for t in state["timeline"]:
        if t["dose"] == "?":
            problems.append(f"{t['no']} 剂量未知：{t['text']}")
        if t["route"] == "?":
            warnings.append(f"{t['no']} 途径缺失：{t['text']}")
        if t["time"] is None:
            warnings.append(f"{t['no']} 没有开始时间")
    rnd = state["attempts"] + 1
    print(f"[质控 第{rnd}轮] 拦下 {len(problems)} 条，提示 {len(warnings)} 条")
    return {"qc": {"passed": not problems, "problems": problems, "warnings": warnings}}


# ---------- 条件边：质控之后的「护士站」 ----------
def route_after_qc(state: VisitState) -> str:
    """返回值是下一个节点的名字——这就是条件边和普通边的全部区别

    判断依据（按顺序）：
      1) qc.passed 为 True → 通过，去 render_report
      2) 重写次数用完（attempts >= MAX_REWRITE）→ 只能带着问题去 render_report
         （真实系统里就是「带着质控报告交医生」，不能永远卡在流程里）
      3) 否则 → 退回 fix_issues 重写一轮
    """
    # 参考答案（TODO 1）：三条规则逐条翻译，顺序就是优先级
    if state["qc"]["passed"]:
        return "render_report"                     # 规则1：通过 → 出报告
    if state["attempts"] >= MAX_REWRITE:
        return "render_report"                     # 规则2：次数用完 → 带着问题出报告
    return "fix_issues"                            # 规则3：否则退回重写


# ---------- 节点 4：重写（被退回后做的事）----------
DOSE_PATTERN = re.compile(r"(\d+(?:\.\d+)?)\s*(ml|mg|g|iu|ug|片|支|袋|粒)", re.IGNORECASE)


def fix_issues(state: VisitState) -> dict:
    """对剂量未知的条目做一轮修复：
       能从医嘱文本里提取到剂量 → 用提取值（依据是原文，合法）；
       提取不到 → 标【待医生补充】（没有依据，绝不编数）"""
    # 注意：复制一份再改，返回新列表。图编排里节点不该偷偷改上游的数据
    timeline = [dict(t) for t in state["timeline"]]
    notes = list(state["notes"])

    fixed, marked = 0, 0
    for t in timeline:
        if t["dose"] != "?":
            continue
        m = DOSE_PATTERN.search(t["text"])
        if m:
            # 参考答案（TODO 2）：数字 + 单位拼成剂量；依据是医嘱原文，所以合法
            t["dose"] = f"{m.group(1)}{m.group(2)}"
            fixed += 1
        else:
            t["dose"] = "【待医生补充】"
            marked += 1
    notes.append(f"重写第{state['attempts'] + 1}轮：文本提取 {fixed} 条，标记待补充 {marked} 条")
    print(f"[重写 第{state['attempts'] + 1}轮] 提取 {fixed} 条，标记 {marked} 条")
    return {"timeline": timeline, "attempts": state["attempts"] + 1, "notes": notes}


# ---------- 节点 5：输出（比第 1 课多了质控摘要）----------
def render_report(state: VisitState) -> dict:
    lines = [f"用药时间线：{state['patient_id']}（共 {len(state['timeline'])} 条）"]
    for t in state["timeline"]:
        time_txt = t["time"] or "时间未记录"
        lines.append(
            f"{t['no']} | {time_txt} | {t['text']} {t['dose']} {t['route']}"
            f" | 出处：{t['source']}"
        )
    lines.append("")
    lines.append(f"质控：检查 {state['attempts'] + 1} 轮，重写 {state['attempts']} 轮，"
                 f"结论 = {'通过' if state['qc']['passed'] else '未完全通过（见下）'}")
    for n in state["notes"]:
        lines.append(f"· {n}")
    for w in state["qc"]["warnings"]:
        lines.append(f"· 提示：{w}")
    for p in state["qc"]["problems"]:
        lines.append(f"· 遗留问题：{p}")
    return {"report": "\n".join(lines)}


# ---------- 接线：这次的图不再是一条直线 ----------
def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("build_timeline", build_timeline)
    g.add_node("qc_check", qc_check)
    g.add_node("fix_issues", fix_issues)
    g.add_node("render_report", render_report)

    g.add_edge(START, "load_orders")
    g.add_edge("load_orders", "build_timeline")
    g.add_edge("build_timeline", "qc_check")

    # 条件边：qc_check 之后走哪条路，由 route_after_qc 的返回值决定
    g.add_conditional_edges("qc_check", route_after_qc)

    # 回头路：重写完必须再送回质控，让 qc_check 重新裁决
    g.add_edge("fix_issues", "qc_check")

    g.add_edge("render_report", END)

    app = g.compile()
    final = app.invoke({
        "patient_id": patient_id, "orders": [], "timeline": [], "report": "",
        "attempts": 0, "qc": {}, "notes": [],
    })

    print()
    print(final["report"])

    out_dir = os.path.join(os.path.dirname(__file__), "..", "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"lesson02_{patient_id}.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(final["report"] + "\n")
    print(f"[已保存] {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
