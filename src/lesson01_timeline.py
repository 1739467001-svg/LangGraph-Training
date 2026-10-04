# -*- coding: utf-8 -*-
"""
第 1 课：LangGraph 三件套 —— State（资料袋）/ Node（科室）/ Edge（箭头）

流程是一条直线图：
    读入医嘱(load_orders) → 整理时间线(build_timeline) → 输出报告(render_report)

运行：
    /usr/bin/python3 src/lesson01_timeline.py 测试患者甲
    （参数是「清洗后的数据/住院/」下的文件夹名；住院号 FAKE0001 在表里，会被节点读出来）
"""
import os
import sys
from datetime import datetime

# 数据安全铁律第 4 条：关闭一切云端追踪（写死在代码最前面）
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import TypedDict
from langgraph.graph import StateGraph, START, END
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")


# ---------- 1) State：全流程唯一的一个「资料袋」 ----------
class VisitState(TypedDict):
    patient_id: str     # 如 FAKE0001
    orders: list        # 从表格读出的原始医嘱，每条一个 dict
    timeline: list      # 编号排序后的时间线
    report: str         # 最终输出的文本


# ---------- 2) Node：每个节点只做一件事 ----------
# 约定：节点函数接收整个资料袋，但【只返回它更新的那几个字段】——
# 就像科室往袋子里塞一张新单子，而不是把整个袋子重抄一遍。
def load_orders(state: VisitState) -> dict:
    """读入该患者的医嘱表，取出需要的列

    注意：参数是文件夹名（测试患者甲），住院号 FAKE0001 在表格里——
    所以这个节点会顺手把表里的 PATIENT_ID 写回资料袋（节点可以中途更新状态）。
    """
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]  # 第 1 行是表头
    idx = {h: i for i, h in enumerate(headers)}

    orders = []
    patient_id = state["patient_id"]
    for row_no, row in enumerate(rows, start=2):          # Excel 第 2 行起是数据
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])      # 用表里的住院号作为正式 ID
        if not text:
            continue
        orders.append({
            "row": row_no,                                # 出处①：第几行
            "order_no": row[idx["ORDER_NO"]],             # 出处②：医嘱号
            "text": str(text),
            # 参考答案（TODO 1）：和取 text 的写法完全一样，只是换了列名
            "dosage": row[idx["DOSAGE"]],
            "units": row[idx["DOSAGE_UNITS"]],
            "route": row[idx["ADMINISTRATION"]],
            "start": row[idx["START_DATE_TIME"]],         # 开始时间（可能为 None）
            "doctor": row[idx["DOCTOR"]],
        })
    # 返回两个要更新的字段：orders 是新算出来的，patient_id 是顺手改写的
    return {"orders": orders, "patient_id": patient_id}


def build_timeline(state: VisitState) -> dict:
    """按开始时间排序，逐条编号 T001、T002……每条都带出处"""
    orders = list(state["orders"])
    orders.sort(key=lambda o: o["start"] or datetime.max)  # 没写开始时间的排最后
    timeline = []
    for i, o in enumerate(orders, start=1):
        dose = " ".join(str(x) for x in (o["dosage"], o["units"]) if x) or "?"
        timeline.append({
            "no": f"T{i:03d}",
            "time": o["start"],
            "text": o["text"],
            "dose": dose,
            "route": o["route"] or "?",
            "source": f"医嘱.xlsx 第{o['row']}行 医嘱号{o['order_no']}",
        })
    return {"timeline": timeline}


def render_report(state: VisitState) -> dict:
    """把时间线拼成最终文本"""
    lines = [f"用药时间线：{state['patient_id']}（共 {len(state['timeline'])} 条）"]
    for t in state["timeline"]:
        # 参考答案（TODO 2）：一行 f-string 拼完。
        # 注意两点：① time 可能为 None，用 or 给兜底文案；② 防编造原则——
        # 出处 t["source"] 是每行天生自带的，不是事后补的
        time_txt = t["time"] or "时间未记录"
        lines.append(
            f"{t['no']} | {time_txt} | {t['text']} {t['dose']} {t['route']}"
            f" | 出处：{t['source']}"
        )
    return {"report": "\n".join(lines)}


# ---------- 3) Edge：把科室挂上图、画上箭头 ----------
def main():
    # 默认用第一位患者；也可以换 测试患者乙 / 丙 / 丁 / 戊
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)      # 挂上第一个科室
    g.add_node("build_timeline", build_timeline)
    g.add_node("render_report", render_report)

    g.add_edge(START, "load_orders")            # 入口箭头
    g.add_edge("load_orders", "build_timeline")
    g.add_edge("build_timeline", "render_report")
    g.add_edge("render_report", END)            # 出口箭头

    app = g.compile()                            # 流程图定稿

    # invoke = 让病人带着初始资料袋走完全程，返回走完后的袋子
    final = app.invoke({"patient_id": patient_id, "orders": [], "timeline": [], "report": ""})

    print(final["report"])

    out_dir = os.path.join(os.path.dirname(__file__), "..", "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"lesson01_{patient_id}.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(final["report"] + "\n")
    print(f"\n[已保存] {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
