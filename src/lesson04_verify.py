# -*- coding: utf-8 -*-
"""
第 4 课：核对节点 —— 用代码逐字核对模型的输出

（防编造原则第 3 条：数值、药名、日期要用代码和出处逐字核对）

核心思想：**模型说的话，一律先当嫌疑人。**
每条用药事实自己报了出处（第N行 医嘱号M），核对就是拿着这个出处回表格查：
  ① 出处存在吗？      第 9999 行 → 当场识破
  ② 药名对得上吗？    模型写下的每个字，都必须能在出处原文里逐字找到
  ③ 剂量对得上吗？    数值等价即可（100 == 100.0，空格无关）
对不上的 → 剔除 + 记录原因。**剔除清单要给医生看，不能悄悄扔**——
这是审计轨迹，也是第 7 课「编造数」指标的原料。

图：START → load_orders → extract_facts → verify_facts → render_draft → END
          （同第 3 课）      （本课新角色）

运行：/usr/bin/python3 src/lesson04_verify.py 测试患者甲
"""
import json
import os
import re
import sys

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import List, TypedDict
from pydantic import BaseModel, Field
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langchain_core.prompts import ChatPromptTemplate
from langgraph.graph import StateGraph, START, END
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")
SHOW_LIMIT = 30


# ---------- 与第 3 课相同：合同 / 替身模型 / 提示词 / 链 ----------
class MedicationFact(BaseModel):
    drug: str = Field(description="药品名称，逐字照抄医嘱原文")
    dose: str = Field(description="剂量+单位；原文找不到依据就填【待医生补充】")
    route: str = Field(description="给药途径")
    source: str = Field(description="出处，格式：医嘱.xlsx 第N行 医嘱号M")
    freq: str = Field(description="用药频次，逐字照抄医嘱原文；找不到依据就填【待医生补充】")


class FactList(BaseModel):
    facts: List[MedicationFact]


FIXED_SCRIPT = [
    {"drug": "0.9%氯化钠注射液", "dose": "100ml", "route": "静脉滴注",
     "freq": "每日1次", "source": "医嘱.xlsx 第20行 医嘱号10"},
    {"drug": "注射用头孢呋辛钠", "dose": "1.5g", "route": "静脉滴注",
     "freq": "每8小时1次", "source": "医嘱.xlsx 第9999行 医嘱号404"},  # 错弹：出处不存在
    {"drug": "阿司匹林肠溶片", "dose": "【待医生补充】", "route": "口服",
     "freq": "每日1次", "source": "医嘱.xlsx 第33行 医嘱号21"},
]


class FakeMedModel(BaseChatModel):
    @property
    def _llm_type(self) -> str:
        return "fake-med-model"

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        content = json.dumps({"facts": FIXED_SCRIPT}, ensure_ascii=False)
        return ChatResult(generations=[ChatGeneration(message=AIMessage(content=content))])


prompt = ChatPromptTemplate.from_messages([
    ("system", """你是一名住院病历的用药信息抽取助手。请从下面的医嘱列表里抽取用药事实。

规则：
1. 只使用医嘱原文里出现的信息，药名、剂量逐字照抄，不改写不翻译；
2. 每一条事实都必须给出处，格式：医嘱.xlsx 第N行 医嘱号M；
3. 原文里没有的信息，一律填【待医生补充】，绝不猜测、绝不编造数值；
4. 你只负责抽取事实，不判断用药对错、不下诊断、不建议换药——那是医生的工作。"""),
    ("human", "医嘱列表（共 {order_count} 条，下面展示前 {shown} 条）：\n{orders_text}"),
])


def parse_facts(ai_msg: AIMessage) -> FactList:
    return FactList.model_validate_json(ai_msg.content)


extraction_chain = prompt | FakeMedModel() | parse_facts


# ---------- 节点：读表（同第 3 课）----------
class VisitState(TypedDict):
    patient_id: str
    orders: list
    facts: list
    verified: list   # 核对通过的事实
    rejected: list   # 剔除的事实（带 reason）
    draft: str


def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        dose = " ".join(str(x) for x in (row[idx["DOSAGE"]], row[idx["DOSAGE_UNITS"]]) if x)
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]],
            "text": str(text), "dose": dose or "【待医生补充】",
            "route": row[idx["ADMINISTRATION"]] or "【待医生补充】",
            "freq": row[idx["FREQUENCY"]] or "【待医生补充】",
        })
    return {"orders": orders, "patient_id": patient_id}


def extract_facts(state: VisitState) -> dict:
    orders = state["orders"]
    shown = orders[:SHOW_LIMIT]
    orders_text = "\n".join(
        f"第{o['row']}行 医嘱号{o['order_no']}：{o['text']} {o['dose']} {o['route']} {o['freq']}"
        for o in shown
    )
    print(f"[抽取] 把 {len(shown)}/{len(orders)} 条医嘱交给模型……")
    result: FactList = extraction_chain.invoke({
        "order_count": len(orders), "shown": len(shown), "orders_text": orders_text,
    })
    facts = [f.model_dump() for f in result.facts]
    print(f"[抽取] 模型交货 {len(facts)} 条")
    return {"facts": facts}


# ---------- 本课主角：核对节点 ----------
def _norm(s) -> str:
    """剂量归一化：去掉空格、大小写统一、100.0 视同 100"""
    s = str(s).replace(" ", "").lower()
    return re.sub(r"(\d+)\.0+(?=\D|$)", r"\1", s)


def check_dose(fact_dose: str, row_dose: str) -> bool:
    """剂量核对：事实剂量是【待医生补充】→ 诚实，不算错；否则必须数值等价"""
    if "待医生补充" in fact_dose:
        return True
    if "待医生补充" in row_dose:
        return False   # 表里都没有的剂量，模型却说有 —— 这就是编造
    return _norm(fact_dose) == _norm(row_dose)


def verify_facts(state: VisitState) -> dict:
    """参考答案：逐字核对三步——出处存在 → 药名在原文里 → 剂量数值等价"""
    orders_by_row = {o["row"]: o for o in state["orders"]}
    verified, rejected = [], []
    for f in state["facts"]:
        m = re.search(r"第(\d+)行", f["source"])
        row = orders_by_row.get(int(m.group(1))) if m else None
        if row is None:
            rejected.append({**f, "reason": f"出处不存在：{f['source']}"})
            continue
        if f["drug"] not in row["text"]:
            rejected.append({**f, "reason": f"药名与出处不符：第{row['row']}行原文是『{row['text']}』"})
            continue
        if not check_dose(f["dose"], row["dose"]):
            rejected.append({**f, "reason": f"剂量与出处不符：第{row['row']}行原文剂量是『{row['dose']}』"})
            continue
        verified.append(f)
    print(f"[核对] 通过 {len(verified)} 条，剔除 {len(rejected)} 条")
    return {"verified": verified, "rejected": rejected}


def render_draft(state: VisitState) -> dict:
    n_v, n_r = len(state["verified"]), len(state["rejected"])
    lines = [f"用药草稿（核对后）：{state['patient_id']}"
             f" —— 模型交货 {n_v + n_r} 条，通过 {n_v} 条，剔除 {n_r} 条"]
    for i, f in enumerate(state["verified"], start=1):
        lines.append(
            f"{i}. {f['drug']} | {f['dose']} | {f['route']} | 频次:{f['freq']}"
            f" | 出处：{f['source']}"
        )
    lines.append("")
    lines.append("剔除清单（要给医生看，不能悄悄扔）：")
    for f in state["rejected"]:
        lines.append(f"· {f['drug']}（{f['dose']}）—— {f['reason']}")
    lines.append("")
    lines.append("⚠ 草稿仅含核对通过的事实；最终须经医生审核确认。")
    return {"draft": "\n".join(lines)}


def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("extract_facts", extract_facts)
    g.add_node("verify_facts", verify_facts)
    g.add_node("render_draft", render_draft)
    g.add_edge(START, "load_orders")
    g.add_edge("load_orders", "extract_facts")
    g.add_edge("extract_facts", "verify_facts")
    g.add_edge("verify_facts", "render_draft")
    g.add_edge("render_draft", END)

    app = g.compile()
    final = app.invoke({"patient_id": patient_id, "orders": [], "facts": [],
                        "verified": [], "rejected": [], "draft": ""})

    print()
    print(final["draft"])

    out_dir = os.path.join(os.path.dirname(__file__), "..", "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"lesson04_{patient_id}.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(final["draft"] + "\n")
    print(f"[已保存] {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
