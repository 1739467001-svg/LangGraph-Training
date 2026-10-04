# -*- coding: utf-8 -*-
"""
第 3 课：LangChain —— 提示词模板、调用大模型、结构化输出

两个库的分工（用比喻记）：
  LangChain = 和大模型打交道的「同声传译套装」：模板帮你把话填好，
              输出解析帮你说好返回格式，拿回来的 JSON 当场验货
  LangGraph = 编排流程的图（前两课学的三件套）

本机没有 API key，所以主角是「替身模型」：不管你问什么，它都背同一份台词。
但它走的链路（模板 → 模型 → 解析 → 校验）和真模型一模一样——
以后接真模型只是换掉一个零件。

图结构：
    START → load_orders → extract_facts(大模型) → render_draft → END

⚠ 本课固定台词里埋了一颗「错弹」（一条出处不存在的用药事实），
  第 4 课的核对节点会抓住它——先别改代码，让它活着。

运行：
    /usr/bin/python3 src/lesson03_llm_extract.py 测试患者甲
"""
import json
import os
import sys

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import List
from pydantic import BaseModel, Field
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langchain_core.prompts import ChatPromptTemplate
from langgraph.graph import StateGraph, START, END
from typing import TypedDict
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")
SHOW_LIMIT = 30  # 给模型看的医嘱条数上限（真模型按 token 计费，也要防上下文爆炸）


# ---------- 1) 结构化输出的「合同」：模型必须按这个格式交货 ----------
class MedicationFact(BaseModel):
    drug: str = Field(description="药品名称，逐字照抄医嘱原文")
    dose: str = Field(description="剂量+单位；原文找不到依据就填【待医生补充】")
    route: str = Field(description="给药途径")
    source: str = Field(description="出处，格式：医嘱.xlsx 第N行 医嘱号M")
    # 参考答案（TODO 2）：频次也是「逐字照抄，没依据就待补充」的同款规矩
    freq: str = Field(description="用药频次，逐字照抄医嘱原文（如「每日1次」）；找不到依据就填【待医生补充】")


class FactList(BaseModel):
    facts: List[MedicationFact]


# ---------- 2) 替身模型：没有 API key 时的「背台词演员」 ----------
FIXED_SCRIPT = [
    {"drug": "0.9%氯化钠注射液", "dose": "100ml", "route": "静脉滴注",
     "freq": "每日1次", "source": "医嘱.xlsx 第20行 医嘱号10"},
    {"drug": "注射用头孢呋辛钠", "dose": "1.5g", "route": "静脉滴注",
     "freq": "每8小时1次", "source": "医嘱.xlsx 第9999行 医嘱号404"},  # 错弹：这条出处不存在
    {"drug": "阿司匹林肠溶片", "dose": "【待医生补充】", "route": "口服",
     "freq": "每日1次", "source": "医嘱.xlsx 第33行 医嘱号21"},
]


class FakeMedModel(BaseChatModel):
    """继承 BaseChatModel，接口和真模型完全一致，只是 _generate 永远背同一份台词"""

    @property
    def _llm_type(self) -> str:
        return "fake-med-model"

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        content = json.dumps({"facts": FIXED_SCRIPT}, ensure_ascii=False)
        return ChatResult(generations=[ChatGeneration(message=AIMessage(content=content))])


def make_llm() -> BaseChatModel:
    # 铁律 6：API key 只从环境变量读；而且「清洗后的数据」能不能发给某个厂商，
    # 要等你明确批准——批准前一律用替身模型
    if os.environ.get("DRCOPILOT_API_KEY"):
        print("[模型] 检测到 DRCOPILOT_API_KEY，但在你明确批准前，本课仍使用替身模型")
    else:
        print("[模型] 未检测到 API key，使用替身模型（背固定台词）")
    return FakeMedModel()


# ---------- 3) 提示词模板：填空题，不是整段字符串硬拼 ----------
prompt = ChatPromptTemplate.from_messages([
    ("system", """你是一名住院病历的用药信息抽取助手。请从下面的医嘱列表里抽取用药事实。

规则：
1. 只使用医嘱原文里出现的信息，药名、剂量逐字照抄，不改写不翻译；
2. 每一条事实都必须给出处，格式：医嘱.xlsx 第N行 医嘱号M；
3. 原文里没有的信息，一律填【待医生补充】，绝不猜测、绝不编造数值；
4. 你只负责抽取事实，不判断用药对错、不下诊断、不建议换药——那是医生的工作。"""),
    ("human", "医嘱列表（共 {order_count} 条，下面展示前 {shown} 条）：\n{orders_text}"),
])


def parse_facts(ai_msg: AIMessage) -> FactList:
    """输出解析+验货：模型回的 JSON 当场用合同校验，不合格直接报错"""
    return FactList.model_validate_json(ai_msg.content)


# LangChain 表达式（LCEL）：用 | 把「模板 → 模型 → 解析」串成一条链
extraction_chain = prompt | make_llm() | parse_facts


# ---------- 4) 把这块积木装进图 ----------
class VisitState(TypedDict):
    patient_id: str
    orders: list
    facts: list    # 模型交货的用药事实（每条都是通过合同校验的 dict）
    draft: str


def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        dose = " ".join(str(x) for x in (row[idx["DOSAGE"]], row[idx["DOSAGE_UNITS"]]) if x)
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]],
            "text": str(text), "dose": dose or "【待医生补充】",
            "route": row[idx["ADMINISTRATION"]] or "【待医生补充】",
            "freq": row[idx["FREQUENCY"]] or "【待医生补充】",
        })
    return {"orders": orders, "patient_id": patient_id}


def extract_facts(state: VisitState) -> dict:
    orders = state["orders"]
    shown = orders[:SHOW_LIMIT]
    orders_text = "\n".join(
        f"第{o['row']}行 医嘱号{o['order_no']}：{o['text']} {o['dose']} {o['route']} {o['freq']}"
        for o in shown
    )
    print(f"[抽取] 把 {len(shown)}/{len(orders)} 条医嘱交给模型……")
    result: FactList = extraction_chain.invoke({
        "order_count": len(orders),
        "shown": len(shown),
        "orders_text": orders_text,
    })
    facts = [f.model_dump() for f in result.facts]
    print(f"[抽取] 模型交货 {len(facts)} 条，全部通过合同校验")
    return {"facts": facts}


def render_draft(state: VisitState) -> dict:
    lines = [f"用药草稿（{state['patient_id']}，模型给出 {len(state['facts'])} 条用药事实）"]
    for i, f in enumerate(state["facts"], start=1):
        lines.append(
            f"{i}. {f['drug']} | {f['dose']} | {f['route']} | 频次:{f.get('freq', '—')}"
            f" | 出处：{f['source']}"
        )
    lines.append("")
    lines.append("⚠ 本草稿由模型起草：每一条的药名、剂量、出处都必须核对后才能使用。")
    return {"draft": "\n".join(lines)}


def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("extract_facts", extract_facts)
    g.add_node("render_draft", render_draft)
    g.add_edge(START, "load_orders")
    g.add_edge("load_orders", "extract_facts")
    g.add_edge("extract_facts", "render_draft")
    g.add_edge("render_draft", END)

    app = g.compile()
    final = app.invoke({"patient_id": patient_id, "orders": [], "facts": [], "draft": ""})

    print()
    print(final["draft"])

    out_dir = os.path.join(os.path.dirname(__file__), "..", "outputs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"lesson03_{patient_id}.txt")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(final["draft"] + "\n")
    print(f"[已保存] {os.path.abspath(out_path)}")


if __name__ == "__main__":
    main()
