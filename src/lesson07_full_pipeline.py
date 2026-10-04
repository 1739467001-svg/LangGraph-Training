# -*- coding: utf-8 -*-
"""
第 7 课（收官）：合成总图 —— 并行抽取 → 逐字核对 → 医生审核 → 回填 XML → 对照评估

把前六课的积木拼成一张图（这就是真实项目 DrCopilot 第一里程碑的骨架）：

    START → load_orders ─┬─ Send×N → extract_chunk ──→ verify_all ──批准→ finalize ─┐
                         │      （并行抽取）      （逐字核对）   ↑reject        │
                         │                        │            │              ↓
                         │                        │重试(Send)  doctor_review → export_eval → END
                         └────────────────────────┘            (interrupt)  ↑驳回→ archive ┘

评估指标（第 7 课的灵魂，真实项目向医院交验收就靠它）：
  标准答案 S = 医生写的 用药记录_*.xml 里的药物集合（一家医院一份药一份记录）
  覆盖率 = 草稿命中标准答案的药物数 ÷ |S|
  多报数 = 草稿里有、标准答案里没有的药物数 → 交给医生判断，不自动定罪
  编造数 = 核对节点剔除的事实数（出处不存在/药名不符/剂量不符）

本课的「替身模型」升级为规则基线抽取器：从喂给它的医嘱文本里用规则挑出
像药的条目——真模型上线前的标准做法：先让确定性基线把整条流水线跑通，
评估框架先转起来，模型换上去只是换零件。

运行：
    /usr/bin/python3 src/lesson07_full_pipeline.py 测试患者甲 批准
"""
import json
import math
import os
import re
import sys
import zipfile
import glob
from datetime import datetime

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import Annotated, List, TypedDict
import operator
from pydantic import BaseModel, Field
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langchain_core.prompts import ChatPromptTemplate
from langgraph.graph import StateGraph, START, END
from langgraph.types import Send, Command, interrupt
from langgraph.checkpoint.memory import InMemorySaver
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA_DIR = os.path.join(ROOT, "清洗后的数据", "住院")
OUT_DIR = os.path.join(ROOT, "outputs")
CHUNK_SIZE = 20
MAX_RETRY = 1
PASS_RATE_FLOOR = 0.5   # 核对通过率低于它就退回重抽（最多 MAX_RETRY 次）


# ---------- 合同 ----------
class MedicationFact(BaseModel):
    drug: str = Field(description="药品名称，逐字照抄医嘱原文")
    dose: str = Field(description="剂量+单位；原文找不到依据就填【待医生补充】")
    route: str = Field(description="给药途径")
    source: str = Field(description="出处，格式：医嘱.xlsx 第N行 医嘱号M")
    freq: str = Field(description="用药频次，逐字照抄医嘱原文；找不到依据就填【待医生补充】")


class FactList(BaseModel):
    facts: List[MedicationFact]


# ---------- 规则基线抽取器（替身模型的升级版）----------
DRUG_HINT = re.compile(r"注射液|注射用|吸入用|片|胶囊|颗粒|口服液|滴丸|软膏|丸|滴眼液")


class BaselineModel(BaseChatModel):
    """确定性基线：从输入的医嘱文本里按规则挑药。
       接口与真模型完全一致；真模型获批后换掉它，全图不用改。"""

    @property
    def _llm_type(self) -> str:
        return "baseline-rule-extractor"

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        text = messages[-1].content   # 模板填好后的人话消息：每行一条医嘱
        facts = []
        for line in text.splitlines():
            if "：" not in line:
                continue
            head, body = line.split("：", 1)
            m_row, m_no = re.search(r"第(\d+)行", head), re.search(r"医嘱号(\d+)", head)
            parts = body.split()
            if not (m_row and m_no and parts):
                continue
            name = parts[0]
            if not DRUG_HINT.search(name):
                continue                     # 不是药的（护理/检验/饮食）直接跳过
            has_dose = len(parts) >= 3 and re.match(r"\d", parts[1])
            facts.append({
                "drug": name,
                "dose": " ".join(parts[1:3]) if has_dose else "【待医生补充】",
                "route": parts[3] if has_dose and len(parts) >= 4 else "【待医生补充】",
                "freq": parts[4] if has_dose and len(parts) >= 5 else "【待医生补充】",
                "source": f"医嘱.xlsx 第{m_row.group(1)}行 医嘱号{m_no.group(1)}",
            })
            # 注意：这里不再设「每工位最多 N 条」的截断上限——
            # 第 8 课毕业运行发现它会静默丢药（工位16前8条全是溶媒，
            # 头孢曲松排第10被截断 → 丁的覆盖率 0%）。量的问题交给切批大小控制
        content = json.dumps({"facts": facts}, ensure_ascii=False)
        return ChatResult(generations=[ChatGeneration(message=AIMessage(content=content))])


prompt = ChatPromptTemplate.from_messages([
    ("system", """你是一名住院病历的用药信息抽取助手。请从下面的医嘱列表里抽取用药事实。

规则：
1. 只使用医嘱原文里出现的信息，药名、剂量逐字照抄，不改写不翻译；
2. 每一条事实都必须给出处，格式：医嘱.xlsx 第N行 医嘱号M；
3. 原文里没有的信息，一律填【待医生补充】，绝不猜测、绝不编造数值；
4. 你只负责抽取事实，不判断用药对错、不下诊断、不建议换药——那是医生的工作。"""),
    ("human", "医嘱列表（共 {order_count} 条）：\n{orders_text}"),
])

extraction_chain = prompt | BaselineModel() | (lambda m: FactList.model_validate_json(m.content))


# ---------- State ----------
class VisitState(TypedDict):
    folder: str        # 数据文件夹名（测试患者甲）——找 zip 用
    patient_id: str    # 住院号（FAKE0001）——表里的正式身份
    orders: list
    facts: Annotated[list, operator.add]   # 并行工位追加（第 5 课）
    verified: list
    rejected: list
    qc: dict                               # {"pass_rate":..., "retry":..., "attempts":...}
    draft: str
    review: dict
    evaluation: str


# ---------- 节点 ----------
def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        dose = " ".join(str(x) for x in (row[idx["DOSAGE"]], row[idx["DOSAGE_UNITS"]]) if x)
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]], "text": str(text),
            "dose": dose or "【待医生补充】",
            "route": row[idx["ADMINISTRATION"]] or "【待医生补充】",
            "freq": row[idx["FREQUENCY"]] or "【待医生补充】",
        })
    return {"orders": orders, "patient_id": patient_id, "folder": state["patient_id"]}


def fan_out(state: VisitState):
    """并行分发（第 5 课）：切成 N 批，开 N 个工位"""
    orders = state["orders"]
    chunks = [orders[i:i + CHUNK_SIZE] for i in range(0, len(orders), CHUNK_SIZE)]
    print(f"[分发] {len(orders)} 条医嘱 → {len(chunks)} 个工位并行抽取")
    return [Send("extract_chunk", {"patient_id": state["patient_id"],
                                   "chunk_no": i, "orders": ch})
            for i, ch in enumerate(chunks)]


def extract_chunk(state: dict) -> dict:
    chunk, no = state["orders"], state.get("chunk_no", "?")
    orders_text = "\n".join(
        f"第{o['row']}行 医嘱号{o['order_no']}：{o['text']} {o['dose']} {o['route']} {o['freq']}"
        for o in chunk
    )
    result: FactList = extraction_chain.invoke({
        "order_count": len(chunk), "orders_text": orders_text,
    })
    facts = [f.model_dump() for f in result.facts]
    print(f"[工位{no}] {len(chunk)} 条 → 交货 {len(facts)} 条")
    return {"facts": facts}


def _norm(s) -> str:
    s = str(s).replace(" ", "").lower()
    return re.sub(r"(\d+)\.0+(?=\D|$)", r"\1", s)


def check_dose(fact_dose: str, row_dose: str) -> bool:
    if "待医生补充" in fact_dose:
        return True
    if "待医生补充" in row_dose:
        return False
    return _norm(fact_dose) == _norm(row_dose)


def verify_all(state: VisitState) -> dict:
    """逐字核对（第 4 课），并对重试轮次带来的重复事实去重（第 2+5 课合体）"""
    orders_by_row = {o["row"]: o for o in state["orders"]}
    seen, verified, rejected = set(), [], []
    for f in state["facts"]:
        key = (f["drug"], f["dose"], f["source"])
        if key in seen:
            continue
        seen.add(key)
        m = re.search(r"第(\d+)行", f["source"])
        row = orders_by_row.get(int(m.group(1))) if m else None
        if row is None:
            rejected.append({**f, "reason": f"出处不存在：{f['source']}"})
            continue
        if f["drug"] not in row["text"]:
            rejected.append({**f, "reason": f"药名与出处不符：第{row['row']}行原文『{row['text']}』"})
            continue
        if not check_dose(f["dose"], row["dose"]):
            rejected.append({**f, "reason": f"剂量与出处不符：原文『{row['dose']}』"})
            continue
        verified.append(f)

    total = len(verified) + len(rejected)
    rate = len(verified) / total if total else 0.0
    attempts = state["qc"].get("attempts", 0)
    retry = rate < PASS_RATE_FLOOR and attempts < MAX_RETRY and total > 0
    if retry:
        print(f"[核对] 第{attempts + 1}轮：通过率 {rate:.0%} 低于 {PASS_RATE_FLOOR:.0%} → 退回重抽")
    else:
        print(f"[核对] 通过 {len(verified)} 条 / 剔除 {len(rejected)} 条（通过率 {rate:.0%}）")
    return {"verified": verified, "rejected": rejected,
            "qc": {"pass_rate": rate, "retry": retry, "attempts": attempts + (1 if retry else 0)}}


def after_verify(state: VisitState):
    """条件边：要重试 → 返回 Send 列表再开一轮工位；否则 → 先拼草稿再去审核"""
    if state["qc"]["retry"]:
        orders = state["orders"]
        chunks = [orders[i:i + CHUNK_SIZE] for i in range(0, len(orders), CHUNK_SIZE)]
        return [Send("extract_chunk", {"patient_id": state["patient_id"],
                                       "chunk_no": i, "orders": ch})
                for i, ch in enumerate(chunks)]
    return "build_draft"


def doctor_review(state: VisitState) -> dict:
    decision = interrupt({"question": "请审核以下用药草稿（批准 / 驳回）", "draft": state["draft"]})
    print(f"[审核] 医生态度：{decision}")
    return {"review": {"decision": str(decision), "reviewer": "医生（终端）"}}


def route_after_review(state: VisitState) -> str:
    return "finalize" if "批准" in state["review"]["decision"] else "archive"


def finalize(state: VisitState) -> dict:
    lines = ["═══ 正式用药记录（已通过医生审核）═══", state["draft"],
             f"审核记录：{state['review']['reviewer']} · 结论 {state['review']['decision']}"]
    return {"draft": "\n".join(lines)}


def archive(state: VisitState) -> dict:
    lines = ["═══ 未过审草稿（已归档）═══", state["draft"],
             f"审核记录：{state['review']['reviewer']} · 结论 {state['review']['decision']}",
             "⚠ 该稿未经医生确认，不得作为正式病历内容使用。"]
    return {"draft": "\n".join(lines)}


def build_draft(state: VisitState) -> str:
    lines = [f"用药草稿：{state['patient_id']}（核对通过 {len(state['verified'])} 条）"]
    for i, f in enumerate(state["verified"], start=1):
        lines.append(f"{i}. {f['drug']} | {f['dose']} | {f['route']} | 频次:{f['freq']}"
                     f" | 出处：{f['source']}")
    if state["rejected"]:
        lines.append("剔除（未进入草稿）：")
        for f in state["rejected"]:
            lines.append(f"  · {f['drug']} —— {f['reason']}")
    return "\n".join(lines)


# ---------- 评估：读标准答案 + 算指标 ----------
def load_standard_drugs(patient_id: str) -> list:
    """标准答案 = 该患者所有 zip 里 用药记录_*.xml 的药名（从文件名解析）"""
    drugs = []
    for zp in sorted(glob.glob(os.path.join(DATA_DIR, patient_id, "*.zip"))):
        z = zipfile.ZipFile(zp)
        for n in z.namelist():
            m = re.match(r"用药记录_(.+?)用药记录_\d+\.xml", n)
            if m:
                drugs.append(m.group(1))
        z.close()
    return drugs


def norm_name(s: str) -> str:
    """参考答案（TODO 1）：去掉全角/半角括号里的补充说明（双阀）和空格。
       医嘱名常带厂家剂型备注，标准答案名没有——不归一化就对不上"""
    return re.sub(r"（[^）]*）|\([^)]*\)|\s", "", str(s))


def match(a: str, b: str) -> bool:
    """归一化后，一方是另一方的前缀/子串即算匹配（医嘱名和记录名常有详略差异）"""
    na, nb = norm_name(a), norm_name(b)
    return na in nb or nb in na


def evaluate(state: VisitState) -> dict:
    std_all = load_standard_drugs(state["folder"])
    std_unique = set(std_all)
    draft = [f["drug"] for f in state["verified"]]
    matched_std = {s for s in std_unique if any(match(d, s) for d in draft)}
    hits = [d for d in draft if any(match(d, s) for s in std_unique)]
    extra = [d for d in draft if d not in hits]

    # 参考答案（TODO 2）：覆盖率 = 被命中的标准答案药物种数 ÷ 标准答案药物种数
    # 两个坑：分母为 0 要守卫；同名多份记录（酮咯酸×2）按 set 去重
    coverage = (len(matched_std) / len(std_unique)) if std_unique else 0.0

    lines = [
        "═══ 评估报告（只含统计，不含病历内容）═══",
        f"标准答案：{len(std_all)} 份用药记录，去重 {len(std_unique)} 种药",
        f"草稿（核对通过）：{len(draft)} 条",
        f"覆盖：命中 {len(hits)} 条，覆盖 {len(matched_std)}/{len(std_unique)} 种 → 覆盖率 {coverage:.0%}",
        f"多报：{len(extra)} 条（草稿有、标准答案没有——交给医生判断，"
        f"常见原因是普通溶媒/辅助用药不单独写用药记录）",
        f"编造：{len(state['rejected'])} 条（核对剔除）",
    ]
    for d in sorted(set(hits)):
        lines.append(f"  ✓ 草稿命中：{d}")
    for s in sorted(std_unique - matched_std):
        lines.append(f"  ✗ 标准答案未覆盖：{s}")
    return {"evaluation": "\n".join(lines)}


def build_draft_node(state: VisitState) -> dict:
    """把核对通过的事实拼成草稿，递给医生审核（第 6 课 interrupt 的材料）"""
    return {"draft": build_draft(state)}


def export_eval(state: VisitState) -> dict:
    """回填 XML（简版模板；医院真实模板是排版 XML，对接时按模板回填字段）+ 评估"""
    status = "已过审" if "批准" in state["review"]["decision"] else "未过审"
    records = []
    for f in state["verified"]:
        records.append(
            f"  <用药记录 patient=\"{state['patient_id']}\" status=\"{status}\">\n"
            f"    <药品>{f['drug']}</药品>\n"
            f"    <剂量>{f['dose']}</剂量>\n"
            f"    <途径>{f['route']}</途径>\n"
            f"    <频次>{f['freq']}</频次>\n"
            f"    <出处>{f['source']}</出处>\n"
            f"    <审核 医生=\"{state['review']['reviewer']}\" 结论=\"{state['review']['decision']}\"/>\n"
            f"  </用药记录>"
        )
    xml = ("<?xml version=\"1.0\" encoding=\"utf-8\"?>\n"
           f"<用药记录草稿 generated=\"{datetime.now().isoformat(timespec='seconds')}\">\n"
           + "\n".join(records) + "\n</用药记录草稿>\n")
    os.makedirs(OUT_DIR, exist_ok=True)
    xml_path = os.path.join(OUT_DIR, f"lesson07_{state['patient_id']}_用药记录.xml")
    with open(xml_path, "w", encoding="utf-8") as fp:
        fp.write(xml)
    print(f"[导出] XML → {xml_path}")
    return evaluate(state)


def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"
    decision = sys.argv[2] if len(sys.argv) > 2 else "批准"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("extract_chunk", extract_chunk)
    g.add_node("verify_all", verify_all)
    g.add_node("build_draft", build_draft_node)
    g.add_node("doctor_review", doctor_review)
    g.add_node("finalize", finalize)
    g.add_node("archive", archive)
    g.add_node("export_eval", export_eval)

    g.add_edge(START, "load_orders")
    g.add_conditional_edges("load_orders", fan_out)      # Send×N
    g.add_edge("extract_chunk", "verify_all")
    g.add_conditional_edges("verify_all", after_verify)  # 重试(Send) 或 拼草稿
    g.add_edge("build_draft", "doctor_review")
    g.add_conditional_edges("doctor_review", route_after_review)
    g.add_edge("finalize", "export_eval")
    g.add_edge("archive", "export_eval")
    g.add_edge("export_eval", END)

    app = g.compile(checkpointer=InMemorySaver())
    config = {"configurable": {"thread_id": f"full-{patient_id}"}}

    result = app.invoke({
        "folder": patient_id, "patient_id": patient_id, "orders": [], "facts": [],
        "verified": [], "rejected": [], "qc": {}, "draft": "", "review": {}, "evaluation": "",
    }, config)

    if "__interrupt__" in result:
        payload = result["__interrupt__"][0].value
        print("[暂停] 流程停在医生审核站：")
        print(payload["draft"])
        final = app.invoke(Command(resume=decision), config)
        print()
        print(final["evaluation"])
    else:
        print("（没有触发暂停？检查 doctor_review）")


if __name__ == "__main__":
    main()
