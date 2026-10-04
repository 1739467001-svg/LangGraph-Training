# -*- coding: utf-8 -*-
"""
第 6 课：人工审核中断与恢复（interrupt / Command）+ 检查点

场景：草稿生成完了，能不能直接变成正式记录？不能——提示词铁律：
**最终必须经过医生审核**。流程走到「医生审核」这一站时**暂停**：
整张图的状态先存档（检查点），等人表态后从暂停点继续。
批准 → 出正式稿；驳回 → 归档为「未过审草稿」，绝不冒充正式。

本课两个新 API（LangGraph 专门为「人在回路」设计的）：

  interrupt(payload)  —— 节点内调用：图在这里暂停，payload 是递给审核人的材料
  Command(resume=值)  —— 审核人表态后从暂停点恢复；interrupt() 的返回值就是它

  检查点（checkpointer）= 游戏存档：没有它，暂停等于失忆，恢复无从谈起。
  本课用 InMemorySaver（存在内存里，教学够用）；真实项目换成数据库存档
  （SqliteSaver / PostgresSaver），医生第二天再批也能接着走。
  存档按 thread_id 区分——同一张图可以同时挂起多位病人的审核。

图：
    START → load_orders → make_draft → doctor_review(interrupt) ─批准→ finalize → END
                                            └驳回→ archive（归档，绝不冒充正式）

运行（decision 可选 批准/驳回，默认批准）：
    /usr/bin/python3 src/lesson06_review.py 测试患者甲 批准
    /usr/bin/python3 src/lesson06_review.py 测试患者甲 驳回
"""
import json
import os
import sys

# 数据安全铁律第 4 条：关闭一切云端追踪
for _k in ("LANGSMITH_TRACING_V2", "LANGCHAIN_TRACING_V2",
           "LANGSMITH_TRACING", "LANGCHAIN_TRACING"):
    os.environ[_k] = "false"

from typing import List, TypedDict
from pydantic import BaseModel, Field
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langchain_core.prompts import ChatPromptTemplate
from langgraph.graph import StateGraph, START, END
from langgraph.types import Command, interrupt
from langgraph.checkpoint.memory import InMemorySaver
from openpyxl import load_workbook

DATA_DIR = os.path.join(os.path.dirname(__file__), "..", "清洗后的数据", "住院")


# ---------- 合同 / 替身模型 / 链（同第 3～5 课，直接搬）----------
class MedicationFact(BaseModel):
    drug: str = Field(description="药品名称，逐字照抄医嘱原文")
    dose: str = Field(description="剂量+单位；原文找不到依据就填【待医生补充】")
    route: str = Field(description="给药途径")
    source: str = Field(description="出处，格式：医嘱.xlsx 第N行 医嘱号M")
    freq: str = Field(description="用药频次，逐字照抄医嘱原文；找不到依据就填【待医生补充】")


class FactList(BaseModel):
    facts: List[MedicationFact]


FIXED_SCRIPT = [
    {"drug": "0.9%氯化钠注射液", "dose": "100ml", "route": "静脉滴注",
     "freq": "每日1次", "source": "医嘱.xlsx 第20行 医嘱号10"},
    {"drug": "注射用头孢呋辛钠", "dose": "1.5g", "route": "静脉滴注",
     "freq": "每8小时1次", "source": "医嘱.xlsx 第33行 医嘱号21"},
]


class FakeMedModel(BaseChatModel):
    @property
    def _llm_type(self) -> str:
        return "fake-med-model"

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        content = json.dumps({"facts": FIXED_SCRIPT}, ensure_ascii=False)
        return ChatResult(generations=[ChatGeneration(message=AIMessage(content=content))])


prompt = ChatPromptTemplate.from_messages([
    ("system", """你是一名住院病历的用药信息抽取助手。请从下面的医嘱列表里抽取用药事实。

规则：
1. 只使用医嘱原文里出现的信息，药名、剂量逐字照抄，不改写不翻译；
2. 每一条事实都必须给出处，格式：医嘱.xlsx 第N行 医嘱号M；
3. 原文里没有的信息，一律填【待医生补充】，绝不猜测、绝不编造数值；
4. 你只负责抽取事实，不判断用药对错、不下诊断、不建议换药——那是医生的工作。"""),
    ("human", "医嘱列表（共 {order_count} 条）：\n{orders_text}"),
])

extraction_chain = prompt | FakeMedModel() | (lambda m: FactList.model_validate_json(m.content))


# ---------- State ----------
class VisitState(TypedDict):
    patient_id: str
    orders: list
    draft: str
    review: dict   # 审核结论：{"decision": ..., "reviewer": ...}


def load_orders(state: VisitState) -> dict:
    path = os.path.join(DATA_DIR, state["patient_id"], "医嘱.xlsx")
    ws = load_workbook(path, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h else "" for h in next(rows)]
    idx = {h: i for i, h in enumerate(headers)}

    orders, patient_id = [], state["patient_id"]
    for row_no, row in enumerate(rows, start=2):
        text = row[idx["ORDER_TEXT"]]
        if patient_id.startswith("测试") and row[idx["PATIENT_ID"]]:
            patient_id = str(row[idx["PATIENT_ID"]])
        if not text:
            continue
        dose = " ".join(str(x) for x in (row[idx["DOSAGE"]], row[idx["DOSAGE_UNITS"]]) if x)
        orders.append({
            "row": row_no, "order_no": row[idx["ORDER_NO"]], "text": str(text),
            "dose": dose or "【待医生补充】", "freq": row[idx["FREQUENCY"]] or "【待医生补充】",
        })
    return {"orders": orders, "patient_id": patient_id}


def make_draft(state: VisitState) -> dict:
    """草稿节点（教学简化：不做并行和核对，第 7 课总图里它们都会回来）"""
    shown = state["orders"][:10]
    orders_text = "\n".join(
        f"第{o['row']}行 医嘱号{o['order_no']}：{o['text']} {o['dose']} {o['freq']}"
        for o in shown
    )
    result: FactList = extraction_chain.invoke({
        "order_count": len(state["orders"]), "orders_text": orders_text,
    })
    facts = [f.model_dump() for f in result.facts]   # pydantic 对象 → dict
    lines = [f"用药草稿：{state['patient_id']}（模型给出 {len(facts)} 条）"]
    for i, f in enumerate(facts, start=1):
        lines.append(f"{i}. {f['drug']} | {f['dose']} | 频次:{f['freq']} | 出处：{f['source']}")
    return {"draft": "\n".join(lines)}


# ---------- 本课主角：审核节点 ----------
def doctor_review(state: VisitState) -> dict:
    """interrupt(payload) 让整张图在这里暂停。

    payload 会出现在暂停结果里，是递给审核人的「材料袋」；
    恢复时传入的值，就是 interrupt() 的返回值——
    也就是说，这个函数会「停在这里，直到拿到审核意见」。
    """
    decision = interrupt({
        "question": "请审核以下用药草稿（批准 / 驳回）",
        "draft": state["draft"],
    })
    print(f"[审核] 医生态度：{decision}")
    return {"review": {"decision": str(decision), "reviewer": "医生（终端）"}}


def route_after_review(state: VisitState) -> str:
    """参考答案（TODO 1）：审核后的分诊——批准出正式稿，驳回归档"""
    if "批准" in state["review"]["decision"]:
        return "finalize"
    return "archive"


def finalize(state: VisitState) -> dict:
    lines = ["═══ 正式用药记录（已通过医生审核）═══", state["draft"],
             f"审核记录：{state['review']['reviewer']} · 结论 {state['review']['decision']}"]
    return {"draft": "\n".join(lines)}


def archive(state: VisitState) -> dict:
    lines = ["═══ 未过审草稿（已归档）═══", state["draft"],
             f"审核记录：{state['review']['reviewer']} · 结论 {state['review']['decision']}",
             "⚠ 该稿未经医生确认，不得作为正式病历内容使用。"]
    return {"draft": "\n".join(lines)}


def main():
    patient_id = sys.argv[1] if len(sys.argv) > 1 else "测试患者甲"
    decision = sys.argv[2] if len(sys.argv) > 2 else "批准"

    g = StateGraph(VisitState)
    g.add_node("load_orders", load_orders)
    g.add_node("make_draft", make_draft)
    g.add_node("doctor_review", doctor_review)
    g.add_node("finalize", finalize)
    g.add_node("archive", archive)

    g.add_edge(START, "load_orders")
    g.add_edge("load_orders", "make_draft")
    g.add_edge("make_draft", "doctor_review")
    g.add_conditional_edges("doctor_review", route_after_review)
    g.add_edge("finalize", END)
    g.add_edge("archive", END)

    # 检查点 = 游戏存档。同一个 thread_id 一份存档；
    # 暂停后能恢复，全靠它记住了袋子当时的全部内容
    app = g.compile(checkpointer=InMemorySaver())
    config = {"configurable": {"thread_id": f"review-{patient_id}"}}

    # 第一段：走到审核站必然暂停（没有人就无法继续）
    result = app.invoke({"patient_id": patient_id, "orders": [], "draft": "", "review": {}}, config)

    if "__interrupt__" in result:
        payload = result["__interrupt__"][0].value
        print("[暂停] 流程停在医生审核站。递给审核人的材料：")
        print(payload["question"])
        print(payload["draft"])

        # 第二段：审核人表态，从暂停点恢复
        # 参考答案（TODO 2）：Command(resume=...) 把医生的意见送回暂停点，
        # interrupt() 的返回值就是它——图从 doctor_review 内部继续往下走
        final = app.invoke(Command(resume=decision), config)

        print()
        print(final["draft"])
    else:
        print("（没有触发暂停？检查 doctor_review 是否调用了 interrupt）")


if __name__ == "__main__":
    main()
